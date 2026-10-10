"""Packing-list export: same pack body, xlsx or csv download."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)

BODY = {
    "container": {
        "name": "20GP",
        "length": 1000,
        "width": 1000,
        "height": 1000,
        "max_weight": 1000,
    },
    "items": [
        {
            "id": "SKU-A",
            "name": "纸箱A",
            "length": 200,
            "width": 100,
            "height": 300,
            "weight": 12.5,
            "qty": 2,
            "allow_rotate": True,
        },
        {
            "id": "HUGE",
            "name": "放不下",
            "length": 5000,
            "width": 100,
            "height": 100,
            "weight": 1,
            "qty": 1,
            "allow_rotate": False,
        },
    ],
}

DETAIL_HEADERS = [
    "SKU",
    "品名",
    "数量",
    "长(mm)",
    "宽(mm)",
    "高(mm)",
    "重量(kg)",
    "放置X(mm)",
    "放置Y(mm)",
    "放置Z(mm)",
    "旋转",
    "柜体积利用率%",
    "总重量(kg)",
    "状态",
]


def test_export_xlsx_packing_list_fields():
    packed = client.post("/api/pack", json=BODY)
    assert packed.status_code == 200
    stats = packed.json()["stats"]

    res = client.post("/api/pack/export", json=BODY)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "attachment" in res.headers.get("content-disposition", "")

    wb = load_workbook(BytesIO(res.content))
    assert "装箱明细" in wb.sheetnames
    ws = wb["装箱明细"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == DETAIL_HEADERS

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    placed = [r for r in rows if r[13] == "已装入"]
    unplaced = [r for r in rows if r[13] == "未装入"]
    assert len(placed) == stats["placed_count"] == 2
    assert {r[0] for r in placed} == {"SKU-A"}
    assert all(r[1] == "纸箱A" and r[2] == 1 for r in placed)
    assert placed[0][7:10] == (0, 0, 0)
    assert placed[0][10] in ("0°", "90°")
    assert placed[0][11] == round(stats["volume_utilization"] * 100, 2)
    assert placed[0][12] == stats["weight_used_kg"]
    assert len(unplaced) == 1
    assert unplaced[0][0] == "HUGE"
    assert unplaced[0][7] is None

    summary = wb["装箱汇总"]
    summary_map = {row[0].value: row[1].value for row in summary.iter_rows(max_row=8)}
    assert summary_map["柜型"] == "20GP"
    assert summary_map["体积利用率%"] == round(stats["volume_utilization"] * 100, 2)
    assert summary_map["总重量(kg)"] == stats["weight_used_kg"]
    assert summary_map["未装入件数"] == stats["unplaced_count"]


def test_export_csv_has_same_columns():
    res = client.post("/api/pack/export?format=csv", json=BODY)
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    text = res.content.decode("utf-8-sig")
    header, *lines = [ln for ln in text.splitlines() if ln.strip()]
    assert header.split(",") == DETAIL_HEADERS
    assert any(line.startswith("SKU-A,") for line in lines)
    assert any("未装入" in line and "HUGE" in line for line in lines)


def test_export_rejects_unknown_format():
    res = client.post("/api/pack/export?format=pdf", json=BODY)
    assert res.status_code == 400


def test_sample_preset_is_20gp_with_skus():
    res = client.get("/api/sample")
    assert res.status_code == 200
    data = res.json()
    assert data["container"]["name"] == "20GP"
    assert data["container"]["length"] == 5898
    assert len(data["items"]) >= 3
    assert {item["id"] for item in data["items"]} >= {"SKU-A", "SKU-B", "SKU-C"}
