"""Turn a pack result into a packing list (xlsx, csv, or pdf).

Does not reimplement placement. Rows are built from PackResponse plus the
original SKU lines so unplaced pieces still have name and dimensions.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import ItemIn, PackRequest, PackResponse

DETAIL_HEADERS: List[str] = [
    "SKU",
    "品名",
    "数量",
    "长(mm)",
    "宽(mm)",
    "高(mm)",
    "重量(kg)",
    "放置X(mm)",
    "放置Y(mm)",
    "放置Z(mm)",
    "旋转",
    "卸载顺序",
    "装载顺序",
    "装箱紧密度%",
    "总重量(kg)",
    "状态",
]


def _sku(piece_id: str, by_id: Dict[str, ItemIn]) -> str:
    if piece_id in by_id:
        return piece_id
    base, sep, suffix = piece_id.rpartition("#")
    if sep and suffix.isdigit() and base in by_id:
        return base
    return piece_id


def detail_rows(req: PackRequest, result: PackResponse) -> List[list]:
    """One physical carton per row. 数量 is 1; container totals repeat on each row."""
    by_id = {item.id: item for item in req.items}
    util = round(result.stats.volume_utilization * 100, 2)
    total_weight = result.stats.weight_used_kg
    rows: List[list] = []

    for placed in result.placements:
        sku = _sku(placed.id, by_id)
        src = by_id.get(sku)
        name = placed.name or (src.name if src and src.name else sku)
        rows.append(
            [
                sku,
                name,
                1,
                placed.length,
                placed.width,
                placed.height,
                placed.weight,
                placed.x,
                placed.y,
                placed.z,
                "90°" if placed.rotated else "0°",
                placed.unload_order,
                placed.load_order,
                util,
                total_weight,
                "已装入",
            ]
        )

    for piece_id in result.unplaced:
        sku = _sku(piece_id, by_id)
        src = by_id.get(sku)
        name = (src.name if src and src.name else None) or sku
        rows.append(
            [
                sku,
                name,
                1,
                src.length if src else None,
                src.width if src else None,
                src.height if src else None,
                src.weight if src else None,
                None,
                None,
                None,
                None,
                None,
                None,
                util,
                total_weight,
                "未装入",
            ]
        )
    return rows


def summary_pairs(req: PackRequest, result: PackResponse) -> List[tuple]:
    stats = result.stats
    container = result.container
    weight_util = (
        round(stats.weight_utilization * 100, 2)
        if stats.weight_utilization is not None
        else None
    )
    pack_size = (
        f"{stats.pack_length_mm:g}×{stats.pack_width_mm:g}×{stats.pack_height_mm:g}"
        if stats.pack_length_mm
        else None
    )
    cog = f"{stats.cog_x_mm:g},{stats.cog_y_mm:g},{stats.cog_z_mm:g}"
    return [
        ("柜型", container.name),
        (
            "内尺寸(mm)",
            f"{container.length:g}×{container.width:g}×{container.height:g}",
        ),
        ("限重(kg)", container.max_weight),
        ("外包装尺寸(mm)", pack_size),
        ("外包装表面积(m²)", stats.pack_surface_m2),
        ("装箱紧密度%", round(stats.volume_utilization * 100, 2)),
        ("重量利用率%", weight_util),
        ("总重量(kg)", stats.weight_used_kg),
        ("重心(mm)", cog),
        ("重心是否在带内", "是" if stats.cog_ok else "否"),
        ("已装入件数", stats.placed_count),
        ("未装入件数", stats.unplaced_count),
    ]


def _sku_rollup(req: PackRequest, result: PackResponse) -> List[list]:
    by_id = {item.id: item for item in req.items}
    placed_count: Dict[str, int] = {}
    placed_weight: Dict[str, float] = {}
    for placed in result.placements:
        sku = _sku(placed.id, by_id)
        placed_count[sku] = placed_count.get(sku, 0) + 1
        placed_weight[sku] = placed_weight.get(sku, 0.0) + placed.weight

    unplaced_count: Dict[str, int] = {}
    for piece_id in result.unplaced:
        sku = _sku(piece_id, by_id)
        unplaced_count[sku] = unplaced_count.get(sku, 0) + 1

    header = [
        "SKU",
        "品名",
        "订单数量",
        "已装数量",
        "未装数量",
        "长(mm)",
        "宽(mm)",
        "高(mm)",
        "单件重量(kg)",
        "已装重量(kg)",
    ]
    body: List[list] = [header]
    seen = set()
    ordered: Sequence[ItemIn] = req.items
    for item in ordered:
        if item.id in seen:
            continue
        seen.add(item.id)
        body.append(
            [
                item.id,
                item.name or item.id,
                item.qty,
                placed_count.get(item.id, 0),
                unplaced_count.get(item.id, 0),
                item.length,
                item.width,
                item.height,
                item.weight,
                round(placed_weight.get(item.id, 0.0), 2),
            ]
        )
    return body


def to_csv(req: PackRequest, result: PackResponse) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(DETAIL_HEADERS)
    for row in detail_rows(req, result):
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(req: PackRequest, result: PackResponse) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "装箱汇总"
    bold = Font(bold=True)
    for key, value in summary_pairs(req, result):
        summary.append([key, value])
    for row in summary.iter_rows(min_row=1, max_row=8, max_col=1):
        row[0].font = bold

    summary.append([])
    for rollup_row in _sku_rollup(req, result):
        summary.append(rollup_row)
    for cell in summary[10]:
        cell.font = bold

    detail = wb.create_sheet("装箱明细")
    detail.append(DETAIL_HEADERS)
    for cell in detail[1]:
        cell.font = bold
    for row in detail_rows(req, result):
        detail.append(row)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    for sheet in (summary, detail):
        for index, column in enumerate(sheet.columns, start=1):
            width = 12
            for cell in column:
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)) + 2, 28))
            sheet.column_dimensions[get_column_letter(index)].width = width

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cjk_font() -> str:
    for candidate in (
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    ):
        if candidate.is_file():
            return str(candidate)
    raise FileNotFoundError("找不到中文字体，无法导出 PDF")


def _pdf_widths(total: float) -> List[float]:
    weights = [
        1.5, 1.7, 0.7, 0.9, 0.9, 0.9, 1.0, 1.0, 1.0, 1.0, 0.7, 0.95, 0.95, 1.15, 1.05, 0.9,
    ]
    scale = total / sum(weights)
    return [weight * scale for weight in weights]


def to_pdf(req: PackRequest, result: PackResponse) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.add_font("hei", fname=_cjk_font())
    pdf.add_page()
    pdf.set_font("hei", size=12)
    pdf.cell(0, 8, f"装箱单 · {req.container.name}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("hei", size=8)
    for key, value in summary_pairs(req, result):
        shown = "" if value is None else str(value)
        pdf.cell(0, 5, f"{key}：{shown}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("hei", size=7)
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    widths = _pdf_widths(usable)

    def draw_header() -> None:
        pdf.set_font("hei", size=7)
        for header, width in zip(DETAIL_HEADERS, widths):
            pdf.cell(width, 6, header, border=1)
        pdf.ln()

    draw_header()
    for row in detail_rows(req, result):
        if pdf.get_y() > pdf.h - 16:
            pdf.add_page()
            draw_header()
        for value, width in zip(row, widths):
            text = "" if value is None else str(value)
            pdf.cell(width, 5, text, border=1)
        pdf.ln()
    return bytes(pdf.output())
