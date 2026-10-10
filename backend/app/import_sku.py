"""Parse SKU lists from Excel / CSV into ItemIn rows."""

from __future__ import annotations

import csv
import io
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from .models import ItemIn

# header aliases (lower) -> field
ALIASES = {
    "id": "id",
    "sku": "id",
    "货号": "id",
    "编码": "id",
    "name": "name",
    "品名": "name",
    "名称": "name",
    "length": "length",
    "长": "length",
    "长(mm)": "length",
    "长度": "length",
    "width": "width",
    "宽": "width",
    "宽(mm)": "width",
    "宽度": "width",
    "height": "height",
    "高": "height",
    "高(mm)": "height",
    "高度": "height",
    "weight": "weight",
    "重量": "weight",
    "重量(kg)": "weight",
    "qty": "qty",
    "数量": "qty",
    "件数": "qty",
    "allow_rotate": "allow_rotate",
    "可旋转": "allow_rotate",
    "max_stack_load_kg": "max_stack_load_kg",
    "承重": "max_stack_load_kg",
    "承重kg": "max_stack_load_kg",
    "承重上限": "max_stack_load_kg",
    "stackable": "stackable",
    "可承重": "stackable",
}


def _norm_header(h: Any) -> str:
    s = str(h or "").strip().lower().replace(" ", "")
    return s


def _map_headers(headers: List[Any]) -> Dict[int, str]:
    mapping: Dict[int, str] = {}
    for i, h in enumerate(headers):
        key = ALIASES.get(_norm_header(h))
        if key:
            mapping[i] = key
    return mapping


def _truthy(v: Any) -> bool:
    if v is None or v == "":
        return True
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    return s not in {"0", "false", "no", "否", "n"}


def _row_to_item(cells: Dict[str, Any], row_no: int) -> Optional[ItemIn]:
    if "id" not in cells and "length" not in cells:
        return None
    sku = str(cells.get("id") or f"SKU-{row_no}").strip()
    if not sku:
        return None
    try:
        length = float(cells["length"])
        width = float(cells["width"])
        height = float(cells["height"])
    except (KeyError, TypeError, ValueError) as e:
        raise ValueError(f"row {row_no}: need length/width/height") from e
    weight = float(cells.get("weight") or 0)
    qty = int(float(cells.get("qty") or 1))
    name = cells.get("name")
    name_s = str(name).strip() if name not in (None, "") else None
    max_load = cells.get("max_stack_load_kg")
    max_load_f = float(max_load) if max_load not in (None, "") else None
    return ItemIn(
        id=sku,
        name=name_s,
        length=length,
        width=width,
        height=height,
        weight=weight,
        qty=max(1, qty),
        allow_rotate=_truthy(cells.get("allow_rotate", True)),
        max_stack_load_kg=max_load_f,
        stackable=_truthy(cells.get("stackable", True)),
    )


def parse_sku_workbook(raw: bytes, filename: str) -> List[ItemIn]:
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return []
        mapping = _map_headers(rows[0])
        if "length" not in mapping.values():
            raise ValueError("CSV needs columns: SKU/id, 长, 宽, 高, 数量…")
        items: List[ItemIn] = []
        for n, row in enumerate(rows[1:], start=2):
            cells = {mapping[i]: row[i] for i in mapping if i < len(row)}
            it = _row_to_item(cells, n)
            if it:
                items.append(it)
        return items

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = list(next(rows_iter))
    except StopIteration:
        return []
    mapping = _map_headers(headers)
    if "length" not in mapping.values():
        raise ValueError("Excel needs columns: SKU/id, 长, 宽, 高, 数量…")
    items = []
    for n, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        cells = {mapping[i]: row[i] for i in mapping if i < len(row)}
        it = _row_to_item(cells, n)
        if it:
            items.append(it)
    return items
